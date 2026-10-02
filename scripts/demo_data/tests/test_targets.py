"""targets.xlsx follows the master contract (1 sheet, long format, 0-100 %) and only PH-004 July misses."""
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import load_workbook


def _rows(out):
    wb = load_workbook(out.path / "targets.xlsx")
    assert wb.sheetnames == ["targets"], "exactly one sheet"
    return list(wb["targets"].iter_rows(values_only=True))


def test_header_and_long_format(out):
    rows = _rows(out)
    assert rows[0] == ("Store ID", "Metric Key", "Period", "Target Value")
    body = rows[1:]
    assert len(body) == 5 * 17 * 2
    assert {r[1] for r in body} == {"net_sales", "gross_margin_pct"}
    assert all(isinstance(r[2], datetime) and r[2].day == 1 and r[2].hour == 0 for r in body)
    per_key = Counter((r[0], r[1], r[2]) for r in body)
    assert set(per_key.values()) == {1}, "one row per store x metric x month"
    assert all(isinstance(r[3], (int, float)) for r in body)


def test_percent_targets_are_on_a_0_to_100_scale(out):
    wb = load_workbook(out.path / "targets.xlsx")
    ws = wb["targets"]
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "gross_margin_pct":
            assert 5 <= row[3].value <= 95, row[3].value
            assert "%" not in row[3].number_format


def _actual_net(out):
    rx = {p["product_id"] for p in out.rows("products") if p["prescription_required"] == "True"}
    txn = {t["transaction_id"]: (t["store_id"], t["transaction_datetime"][:7]) for t in out.rows("transactions")}
    net = defaultdict(float)
    for li in out.rows("transaction_items"):
        if li["product_id"] not in rx:
            net[txn[li["transaction_id"]]] += float(li["line_total"])
    return net


def test_only_ph004_july_misses_net_sales_target(out):
    net = _actual_net(out)
    misses = []
    for store, key, period, value in _rows(out)[1:]:
        month = period.strftime("%Y-%m")
        if key == "net_sales" and month <= "2026-07" and net[(store, month)] < value:
            misses.append((store, month))
    assert misses == [("PH-004", "2026-07")]
    ph004 = {(p.strftime("%Y-%m")): v for s, k, p, v in _rows(out)[1:] if s == "PH-004" and k == "net_sales"}
    assert net[("PH-004", "2026-07")] <= 0.70 * ph004["2026-07"], "FND-008: July must miss by >= 30%"
