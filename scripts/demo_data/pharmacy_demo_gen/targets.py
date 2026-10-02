"""Store KPI targets per month for the 'Dữ liệu bảng' upload demo.
Master contract: ONE sheet, LONG format, columns exactly `Store ID | Metric Key | Period | Target Value`;
Metric Key in {net_sales, gross_margin_pct}; Period = first-of-month Excel date; percent on a 0-100 scale."""
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from . import config as C
from .tables import Table, parse_dt

HEADER = ("Store ID", "Metric Key", "Period", "Target Value")
METRIC_KEYS = ("net_sales", "gross_margin_pct")
TARGET_MONTHS = [(2025, m) for m in range(8, 13)] + [(2026, m) for m in range(1, 13)]
FIXED_DOC_TIME = datetime(2026, 8, 1, 0, 0, 0)


def retail_actuals(tables: dict[str, Table]) -> dict[tuple[str, tuple[int, int]], dict[str, float]]:
    """Retail net sales (= line_total, excl. tax and Rx lines) and gross margin % (0-100) per store-month."""
    rx = {p["product_id"] for p in tables["products"].rows if p["prescription_required"] == "True"}
    txn = {t["transaction_id"]: (t["store_id"], parse_dt(t["transaction_datetime"]))
           for t in tables["transactions"].rows}
    acc = defaultdict(lambda: {"net": 0.0, "cost": 0.0})
    for li in tables["transaction_items"].rows:
        if li["product_id"] in rx or li["transaction_id"] not in txn:
            continue
        store, dt = txn[li["transaction_id"]]
        a = acc[(store, (dt.year, dt.month))]
        a["net"] += float(li["line_total"])
        a["cost"] += int(li["quantity"]) * float(li["unit_cost"])
    return {k: {"net_sales": v["net"], "gm_pct": (v["net"] - v["cost"]) / v["net"] * 100} for k, v in acc.items()}


def _store_month_targets(actual, store: str, ym: tuple[int, int], recent_gm: float) -> tuple[float, float]:
    if (store, ym) in actual and ym <= (C.HISTORY_END.year, C.HISTORY_END.month):
        net = round(actual[(store, ym)]["net_sales"] * 0.98, -2)       # achieved
        gm = actual[(store, ym)]["gm_pct"] - 0.5                         # achieved, 0-100 scale
    else:
        net = round(actual[(store, (ym[0] - 1, ym[1]))]["net_sales"] * 1.03, -2)
        gm = recent_gm - 0.5
    if (store, ym) == ("PH-004", (2026, 7)):
        # FND-008: July target follows June's level, so the July sales decline misses the target.
        net = round(actual[("PH-004", (2026, 6))]["net_sales"] * C.MONTH_WEIGHT[7] / C.MONTH_WEIGHT[6], -2)
    return float(net), round(gm - 0.05, 1)   # round DOWN to 1 decimal so the target stays below actual


def build_targets(tables: dict[str, Table]) -> list[dict]:
    """One row per store x month x metric (long format)."""
    actual = retail_actuals(tables)
    stores = sorted(s["store_id"] for s in tables["stores"].rows)
    rows = []
    for store in stores:
        recent_gm = sum(actual[(store, (2026, m))]["gm_pct"] for m in (5, 6, 7)) / 3
        for ym in TARGET_MONTHS:
            net, gm = _store_month_targets(actual, store, ym, recent_gm)
            period = datetime(ym[0], ym[1], 1)
            rows.append({"Store ID": store, "Metric Key": "net_sales", "Period": period, "Target Value": net})
            rows.append({"Store ID": store, "Metric Key": "gross_margin_pct", "Period": period, "Target Value": gm})
    return rows


def write_targets_xlsx(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "targets"
    ws.append(list(HEADER))
    for r in rows:
        ws.append([r[h] for h in HEADER])
        ws.cell(row=ws.max_row, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=ws.max_row, column=4).number_format = "0.0" if r["Metric Key"] == "gross_margin_pct" else "0"
    wb.properties.creator = "pharmacy_demo_gen"
    wb.properties.created = FIXED_DOC_TIME
    wb.properties.modified = FIXED_DOC_TIME
    wb.save(path)
