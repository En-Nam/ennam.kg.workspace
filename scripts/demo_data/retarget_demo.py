"""Rebuild the store targets (targets.xlsx -> DAAB file dataset `store_targets`) for a date-shifted demo package.

    uv run python retarget_demo.py --data ../../demo_mockdata_shifted_70 [--as-of 2026-10-09]
                                   [--original-targets ../../demo_mockdata/targets.xlsx] [--output <data>/targets.xlsx]

Why: RET-SALES-001 and RET-MARGIN-001 compare MONTH-TO-DATE actuals with the monthly target prorated by
elapsedDays/daysInMonth (LAAM src/lib/intel/engine/rules.ts, compare: target). A uniform date shift keeps every
rolling-window rule intact but moves the as-of into another calendar month, so the original targets no longer
describe the planted FND-008 (PH-004 misses its target; every other store makes it). The targets are demo
configuration; this script derives them from the shifted data with the formulas below, deterministically.

Actuals use the definition of the DAAB metrics retail.net_sales / retail.gross_margin_pct over the pharmacy_demo
mapping: completed sales, non-Rx products (products.prescription_required = False), net = qty * unit_price -
discount_amount, cost = qty * COALESCE(line unit_cost, product unit_cost).

Formulas (A = as-of, M = calendar month, run_rate(M) = actual(M) * daysInMonth / days of M inside the data window):
  * months before A's month : net = floor100(run_rate * 0.98)             (as the generator: "achieved")
  * A's month, other stores : net = floor100(run_rate * 0.98)             -> MTD ratio >= 1/0.98 = 1.0204
  * A's month, PH-004       : net = round100(run_rate / plant_ratio)      -> MTD ratio ~= plant_ratio
        plant_ratio = original PH-004 July-2026 actual / original July-2026 target (from the unshifted package)
  * months after A's month  : net = round100(actual(same month a year earlier) * 1.03), else run_rate(A) * 1.03
  * gross_margin_pct        : round(actual_gm - 0.5 - 0.05, 1) for months with data, else mean gm of the last
                              3 complete months - 0.55 (same rule as the generator; RET-MARGIN-001 never fires)
Months covered: first month with data .. December of A's year (the generator also covered the year-end).
"""
import argparse
import csv
import json
import math
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from pharmacy_demo_gen.targets import write_targets_xlsx

import shift_dates as S

ACHIEVED_FACTOR = 0.98
GROWTH_FACTOR = 1.03
GM_OFFSET = 0.5
PLANT_STORE = "PH-004"
ORIGINAL_AS_OF = date(2026, 7, 31)   # last day of the unshifted package (Friday)

Daily = dict[tuple[str, date], tuple[float, float]]   # (store, day) -> (net sales, cost)


def month_start(d: date) -> date:
    return d.replace(day=1)


def next_month(m: date) -> date:
    return date(m.year + (m.month == 12), m.month % 12 + 1, 1)


def days_in_month(m: date) -> int:
    return (next_month(m) - m).days


def floor100(x: float) -> float:
    return float(math.floor(x / 100) * 100)


def read_daily(data_dir: Path) -> Daily:
    """Per store and day: front-store (non-Rx) net sales and cost of completed sales."""
    with (data_dir / "products.csv").open(newline="", encoding="utf-8") as fh:
        products = {p["product_id"]: p for p in csv.DictReader(fh)}
    with (data_dir / "transactions.csv").open(newline="", encoding="utf-8") as fh:
        sales = {t["transaction_id"]: (t["store_id"], date.fromisoformat(t["transaction_datetime"][:10]))
                 for t in csv.DictReader(fh) if t["status"].lower() == "completed"}
    acc: dict[tuple[str, date], list[float]] = defaultdict(lambda: [0.0, 0.0])
    with (data_dir / "transaction_items.csv").open(newline="", encoding="utf-8") as fh:
        for li in csv.DictReader(fh):
            p = products.get(li["product_id"])
            if li["transaction_id"] not in sales or (p and p["prescription_required"] == "True"):
                continue
            qty = int(li["quantity"])
            a = acc[sales[li["transaction_id"]]]
            a[0] += qty * float(li["unit_price"]) - float(li["discount_amount"] or 0)
            a[1] += qty * float(li["unit_cost"] or (p or {}).get("unit_cost") or 0)
    return {k: (v[0], v[1]) for k, v in acc.items()}


def _monthly(daily: Daily, as_of: date):
    """(store, month) -> [net, cost, days with data in the window]; window = first data day .. as_of."""
    first = min(d for _, d in daily)
    out: dict[tuple[str, date], list[float]] = defaultdict(lambda: [0.0, 0.0, 0])
    for (s, d), (net, cost) in daily.items():
        if d <= as_of:
            a = out[(s, month_start(d))]
            a[0] += net
            a[1] += cost
    for (s, m), a in out.items():
        lo, hi = max(m, first), min(next_month(m).toordinal() - 1, as_of.toordinal())
        a[2] = hi - lo.toordinal() + 1
    return out, first


def _gm(net: float, cost: float) -> float:
    return (net - cost) / net * 100


def build_targets(daily: Daily, as_of: date, plant_store: str = PLANT_STORE, plant_ratio: float = 0.5666) -> list[dict]:
    monthly, first = _monthly(daily, as_of)
    stores = sorted({s for s, _ in daily})
    a_month = month_start(as_of)
    months = []
    m = month_start(first)
    while m <= date(as_of.year, 12, 1):
        months.append(m)
        m = next_month(m)
    rows = []
    for store in stores:
        complete = [mm for mm in months if mm < a_month and monthly.get((store, mm), [0, 0, 0])[2] == days_in_month(mm)]
        recent_gm = sum(_gm(*monthly[(store, mm)][:2]) for mm in complete[-3:]) / len(complete[-3:])
        for mm in months:
            net, gm = _store_month(monthly, store, mm, a_month, plant_store, plant_ratio, recent_gm)
            period = datetime(mm.year, mm.month, 1)
            rows.append({"Store ID": store, "Metric Key": "net_sales", "Period": period, "Target Value": net})
            rows.append({"Store ID": store, "Metric Key": "gross_margin_pct", "Period": period, "Target Value": gm})
    return rows


def _store_month(monthly, store, m, a_month, plant_store, plant_ratio, recent_gm) -> tuple[float, float]:
    got = monthly.get((store, m))
    if m <= a_month and got and got[0] > 0:
        run_rate = got[0] * days_in_month(m) / got[2]
        if m == a_month and store == plant_store:
            net = float(round(run_rate / plant_ratio, -2))
        else:
            net = floor100(run_rate * ACHIEVED_FACTOR)
        return net, round(_gm(got[0], got[1]) - GM_OFFSET - 0.05, 1)
    prev = monthly.get((store, date(m.year - 1, m.month, 1)))
    if prev and prev[2] == days_in_month(m):
        net = float(round(prev[0] * GROWTH_FACTOR, -2))
    else:
        cur = monthly[(store, a_month)]
        net = float(round(cur[0] * days_in_month(a_month) / cur[2] * GROWTH_FACTOR, -2))
    return net, round(recent_gm - GM_OFFSET - 0.05, 1)


def read_targets(path: Path) -> dict[tuple[str, str, date], float]:
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))[1:]
    return {(st, mk, per.date() if isinstance(per, datetime) else per): float(v) for st, mk, per, v in rows}


def plant_ratio_from_original(daily: Daily, shift_days: int, original_targets, plant_store: str = PLANT_STORE,
                              original_as_of: date = ORIGINAL_AS_OF) -> float:
    """Original month-to-date ratio of the plant store at the original as-of, read from the shifted data."""
    m = month_start(original_as_of)
    lo, hi = m.toordinal() + shift_days, original_as_of.toordinal() + shift_days
    mtd = sum(v[0] for (s, d), v in daily.items() if s == plant_store and lo <= d.toordinal() <= hi)
    target = original_targets[(plant_store, "net_sales", m)]
    return mtd / (target * original_as_of.day / days_in_month(m))


def mtd_ratios(daily: Daily, rows: list[dict], as_of: date) -> dict[str, float]:
    m = month_start(as_of)
    tgt = {r["Store ID"]: r["Target Value"] for r in rows if r["Metric Key"] == "net_sales" and r["Period"].date() == m}
    out = {}
    for store, t in sorted(tgt.items()):
        mtd = sum(v[0] for (s, d), v in daily.items() if s == store and m <= d <= as_of)
        out[store] = round(mtd / (t * as_of.day / days_in_month(m)), 4)
    return out


def retarget(data_dir: Path, as_of: date | None, original_targets: Path, output: Path | None = None) -> dict:
    manifest = json.loads((data_dir / S.MANIFEST).read_text(encoding="utf-8"))
    shift = int(manifest["days"])
    daily = read_daily(data_dir)
    as_of = as_of or max(d for _, d in daily)
    ratio = plant_ratio_from_original(daily, shift, read_targets(original_targets))
    rows = build_targets(daily, as_of, PLANT_STORE, ratio)
    write_targets_xlsx(output or data_dir / "targets.xlsx", rows)
    return {"as_of": as_of.isoformat(), "shift_days": shift, "plant_store": PLANT_STORE,
            "plant_ratio": round(ratio, 4), "as_of_month_ratios": mtd_ratios(daily, rows, as_of), "rows": len(rows)}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", type=Path, required=True, help="shifted package (has SHIFT_MANIFEST.json)")
    ap.add_argument("--as-of", type=date.fromisoformat, default=None, help="default: last sales day of --data")
    ap.add_argument("--original-targets", type=Path, default=S.DEFAULT_INPUT / "targets.xlsx")
    ap.add_argument("--output", type=Path, default=None, help="default: <data>/targets.xlsx")
    args = ap.parse_args(argv)
    if not (args.data / S.MANIFEST).exists():
        print(f"ERROR: {args.data} has no {S.MANIFEST} — run shift_dates.py first", file=sys.stderr)
        return 1
    json.dump(retarget(args.data, args.as_of, args.original_targets, args.output), sys.stdout, indent=2)
    print()
    return 0


if __name__ == "__main__":
    sys.exit(main())
